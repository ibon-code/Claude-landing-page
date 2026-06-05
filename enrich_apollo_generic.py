"""
Apollo.io email enrichment — generic version.
Usage:
    python3 enrich_apollo_generic.py <excel_file> <api_key> [col_company] [col_website] [col_email_out]

Defaults: company=A, website=B, email output=C, name output=D, title output=E
Columns can be given as letters (A, B, C…) or 1-based numbers.
"""

import sys
import re
import time
import requests
import openpyxl

# ── Args ──────────────────────────────────────────────────────────────────────
if len(sys.argv) < 3:
    print("Usage: python3 enrich_apollo_generic.py <excel_file> <api_key> [col_company] [col_website] [col_email_out]")
    sys.exit(1)

FILE_PATH   = sys.argv[1]
API_KEY     = sys.argv[2]
DELAY       = 1.5
SAVE_EVERY  = 50

def col_to_num(c):
    """Accept 'A'/'B'/'C' or '1'/'2'/'3', return 1-based int."""
    c = str(c).strip().upper()
    if c.isalpha():
        return sum((ord(ch) - 64) * (26 ** i) for i, ch in enumerate(reversed(c)))
    return int(c)

COL_COMPANY  = col_to_num(sys.argv[3]) if len(sys.argv) > 3 else 1
COL_WEBSITE  = col_to_num(sys.argv[4]) if len(sys.argv) > 4 else 2
COL_EMAIL    = col_to_num(sys.argv[5]) if len(sys.argv) > 5 else 3
COL_NAME     = COL_EMAIL + 1
COL_TITLE    = COL_EMAIL + 2

SEARCH_URL = "https://api.apollo.io/v1/mixed_people/api_search"
MATCH_URL  = "https://api.apollo.io/v1/people/match"
HEADERS    = {
    "x-api-key":    API_KEY,
    "Content-Type": "application/json",
    "Accept":       "application/json",
}
PRIORITY_TITLES = [
    ["CEO", "Chief Executive Officer", "Founder", "Co-Founder", "Cofounder",
     "Owner", "President", "Managing Director", "General Manager"],
    ["COO", "CTO", "CFO", "CPO", "VP", "Vice President",
     "Head of", "Director", "Partner", "Principal"],
]

# ── Helpers ───────────────────────────────────────────────────────────────────

def extract_domain(cell_value):
    if not cell_value:
        return None
    s = str(cell_value).strip()
    m = re.search(r'HYPERLINK\("([^"]+)"', s, re.IGNORECASE)
    if m:
        s = m.group(1)
    s = re.sub(r'^https?://', '', s)
    s = re.sub(r'^www\.', '', s)
    return s.split('/')[0].split('?')[0].strip() or None

def title_rank(title):
    if not title:
        return 99
    t = title.lower()
    for rank, group in enumerate(PRIORITY_TITLES):
        if any(kw.lower() in t for kw in group):
            return rank
    return 99

def apollo_request(url, payload, retries=2):
    for attempt in range(retries + 1):
        try:
            r = requests.post(url, json=payload, headers=HEADERS, timeout=20)
        except requests.RequestException as e:
            print(f"    ⚠ Network error: {e}")
            return {}
        if r.status_code == 429:
            wait = 60 if attempt == 0 else 120
            print(f"    ⚠ Rate limited — sleeping {wait}s …")
            time.sleep(wait)
            continue
        if r.status_code != 200:
            print(f"    ⚠ Apollo {r.status_code}: {r.text[:120]}")
            return {}
        return r.json()
    return {}

def enrich_company(company, domain):
    data = apollo_request(SEARCH_URL, {
        "page": 1, "per_page": 10,
        "q_organization_domains": domain,
    })
    people = data.get("people") or []
    candidates = [p for p in people if p.get("has_email")] or people
    if not candidates:
        return "", "", ""
    best = min(candidates, key=lambda p: title_rank(p.get("title", "")))
    data2 = apollo_request(MATCH_URL, {"id": best["id"]})
    person = data2.get("person") or {}
    email = person.get("email") or ""
    name  = person.get("name") or \
            f"{person.get('first_name','')} {person.get('last_name','')}".strip()
    title = person.get("title") or person.get("headline") or ""
    return email, name, title

# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    wb = openpyxl.load_workbook(FILE_PATH)
    ws = wb.active

    # Write output headers if empty
    from openpyxl.utils import get_column_letter
    if not ws.cell(1, COL_EMAIL).value:
        ws.cell(1, COL_EMAIL).value = "Email"
    if not ws.cell(1, COL_NAME).value:
        ws.cell(1, COL_NAME).value = "Full Name"
    if not ws.cell(1, COL_TITLE).value:
        ws.cell(1, COL_TITLE).value = "Job Title"

    total = ws.max_row - 1
    found = skipped = no_result = 0

    print(f"File:    {FILE_PATH}")
    print(f"Columns: company={get_column_letter(COL_COMPANY)}, website={get_column_letter(COL_WEBSITE)}, output starts at {get_column_letter(COL_EMAIL)}")
    print(f"Rows to process: {total}")
    print("─" * 55)

    for row_idx in range(2, ws.max_row + 1):
        company        = str(ws.cell(row_idx, COL_COMPANY).value or "").strip()
        domain         = extract_domain(ws.cell(row_idx, COL_WEBSITE).value)
        existing_email = ws.cell(row_idx, COL_EMAIL).value
        progress       = f"Row {row_idx - 1}/{total}"

        if existing_email and str(existing_email).strip():
            print(f"{progress} — {company} — already enriched, skipping")
            skipped += 1
            continue

        if not company and not domain:
            continue

        email, name, title = enrich_company(company, domain or "")

        if email:
            ws.cell(row_idx, COL_EMAIL).value  = email
            ws.cell(row_idx, COL_NAME).value   = name
            ws.cell(row_idx, COL_TITLE).value  = title
            found += 1
            print(f"{progress} — {company} — ✓ {email} ({name}, {title})")
        else:
            no_result += 1
            print(f"{progress} — {company} — no email found")

        if (row_idx - 1) % SAVE_EVERY == 0:
            wb.save(FILE_PATH)
            print(f"  ── checkpoint saved ({row_idx - 1}/{total} rows) ──")

        time.sleep(DELAY)

    wb.save(FILE_PATH)
    print(f"\n{'─' * 55}")
    print(f"Done.  Found: {found}  |  No result: {no_result}  |  Skipped: {skipped}")
    print(f"File saved: {FILE_PATH}")

if __name__ == "__main__":
    main()
