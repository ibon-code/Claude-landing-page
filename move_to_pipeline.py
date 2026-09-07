"""
Moves any row checked "TRUE" in the "Move to Pipeline?" column (sheet
"companies17-08-2026 (1)") into the "Pipeline" sheet, then removes it from
the sourcing sheet. Meant to run frequently (see n8n workflow) so a checked
box gets picked up promptly without the user having to ask.
"""

import openpyxl
from openpyxl.styles import Font

FILE_PATH = "/Users/iacopobon/Desktop/Claude code experiment/Apollo enrich/Nuove liste/Sourcing Summer challange.xlsx"
SOURCE_SHEET = "companies17-08-2026 (1)"
PIPELINE_SHEET = "Pipeline"
CHECKBOX_COL = "Move to Pipeline?"


def is_checked(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().upper() == "TRUE"
    return False


def main():
    wb = openpyxl.load_workbook(FILE_PATH)
    src = wb[SOURCE_SHEET]
    pipe = wb[PIPELINE_SHEET]

    src_headers = [c.value for c in src[1]]
    src_col = {h: i + 1 for i, h in enumerate(src_headers) if h}

    pipe_headers = [c.value for c in pipe[1]]
    pipe_col = {h: i + 1 for i, h in enumerate(pipe_headers) if h}

    rows_to_move = [
        row for row in range(2, src.max_row + 1)
        if is_checked(src.cell(row, src_col[CHECKBOX_COL]).value)
    ]

    if not rows_to_move:
        print("No rows checked — nothing to move.")
        return

    print(f"Moving {len(rows_to_move)} row(s) to Pipeline…")

    next_pipe_row = pipe.max_row + 1
    for row in rows_to_move:
        name = src.cell(row, src_col["Organization Name"]).value
        website = src.cell(row, src_col["Website"]).value

        pipe.cell(next_pipe_row, pipe_col["Startup name "]).value = name
        website_cell = pipe.cell(next_pipe_row, pipe_col["Website"])
        website_cell.value = website
        if website:
            url_str = str(website).strip()
            if not url_str.startswith("http"):
                url_str = "https://" + url_str
            website_cell.hyperlink = url_str
            website_cell.font = Font(color="0563C1", underline="single")
        next_pipe_row += 1
        print(f"  + {name} -> Pipeline")

    for row in sorted(rows_to_move, reverse=True):
        src.delete_rows(row, 1)

    wb.save(FILE_PATH)
    print(f"Done. Pipeline now has {next_pipe_row - 2} startups.")


if __name__ == "__main__":
    main()
