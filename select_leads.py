"""
Weekly select step — scores unscored cached candidates (see harvest_leads.py)
against the investor's criteria with Claude, and appends the best-qualifying
ones to the sourcing sheet.

Quality over quantity: MAX_NEW_LEADS is a ceiling, not a target. If fewer than
MAX_NEW_LEADS candidates clear MIN_MOAT_SCORE this week, only those are added —
weak/replicable "vibe coded" candidates are never padded in just to hit a count.
"""

import json
import os
import re
import sys

import openpyxl
import requests

FILE_PATH = "/Users/iacopobon/Desktop/Claude code experiment/Apollo enrich/Nuove liste/Sourcing Summer challange.xlsx"
SHEET_NAME = "companies17-08-2026 (1)"
CACHE_PATH = os.path.join(os.path.dirname(__file__), "leads_cache.json")

MAX_NEW_LEADS = 10
MIN_MOAT_SCORE = 3  # 1-5 scale; below this it's considered too thin/replicable

ANTHROPIC_KEY = os.environ.get("ANTHROPIC_API_KEY")
if not ANTHROPIC_KEY:
    sys.exit("Set the ANTHROPIC_API_KEY environment variable before running this script.")

CLAUDE_MODEL = "claude-haiku-4-5-20251001"

CRITERIA = """You are screening early-stage startups for a B2B-focused VC investor. Criteria, in order of importance:

1. MOST IMPORTANT: the end customer must plausibly be large corporates / enterprise organizations —
   not consumers, not solo freelancers/creators, not small local businesses, not "prosumer" tools.
   Think: a Fortune 500 / large enterprise procurement or ops team would realistically buy this.
   If the target customer looks like SMBs, indie hackers, or individuals, REJECT regardless of anything else.

2. REJECT anything in the "AI agent" category — AI agent platforms, agent frameworks, browser-automation
   agents, autonomous-agent tooling, or generic LLM/GPT wrapper products. This category is explicitly
   excluded no matter how good it looks otherwise.

2b. REJECT aerospace, space, satellite, and drone-based startups entirely — not an area of interest for
   this investor, regardless of moat or enterprise fit.

3. A tangible link to the physical world (hardware, robotics, sensors, IoT, industrial/manufacturing
   equipment, energy infrastructure, materials science, logistics/physical operations) is a BONUS, not a
   requirement. Pure software is completely fine — do NOT reject or penalize a candidate just for being
   software-only. What matters is criteria 1 (enterprise customer) and 4 (real moat), full stop.

4. Must show a genuinely defensible moat: proprietary data/tech, deep domain complexity, regulatory or
   compliance edge, hardware, network effects, or real technical depth. REJECT anything that looks like a
   thin wrapper or generic CRUD app replicable in about a week of "vibe coding". Be strict — when in doubt, reject.

5. Sector-agnostic beyond the above. Geography: Europe, UK, US, or Middle East (if unclear, lower confidence).
   Very early stage (pre-seed/seed) is assumed — each candidate includes its source and founding info;
   trust that pre-filtering (BetaList = launching now, Apollo = founded within the last ~18 months) already
   handles stage, you don't need to re-derive it.
"""


def load_cache():
    if not os.path.exists(CACHE_PATH):
        sys.exit(f"No cache found at {CACHE_PATH} — run harvest_leads.py first.")
    with open(CACHE_PATH) as f:
        return json.load(f)


def save_cache(cache):
    with open(CACHE_PATH, "w") as f:
        json.dump(cache, f, indent=2, ensure_ascii=False)


def score_candidates(candidates):
    payload_items = [
        {
            "name": c["name"],
            "tagline": c.get("tagline", ""),
            "description": c["description"],
            "website": c["website"],
            "source": c.get("source", "BetaList"),
            "founded_year": c.get("founded_year"),
            "funding_usd": c.get("funding_usd"),
        }
        for c in candidates
    ]

    prompt = (
        CRITERIA
        + "\n\nHere are the candidate startups as a JSON array:\n"
        + json.dumps(payload_items, ensure_ascii=False)
        + "\n\nReturn ONLY a JSON array (no prose, no markdown fences), one object per candidate, in this exact shape:\n"
        + '[{"name": "...", "company_name": "...", "pass": true|false, "enterprise_fit": true|false, "is_ai_agent": true|false, '
        + '"has_physical_component": true|false, "moat_score": 1-5, "geography_guess": "...", "reasoning": "one short sentence covering enterprise fit, physical component, and moat"}]\n'
        + '"name" must be echoed back EXACTLY as given in the input so results can be matched. "company_name" is the '
        + "clean startup name only (e.g. from a news headline like \"Acme raises $2M\", company_name is \"Acme\") — "
        + "strip any funding/news phrasing.\n"
        + "moat_score reflects how defensible the product looks (5 = strong moat, 1 = trivially replicable). "
        + "pass must be false whenever enterprise_fit is false or is_ai_agent is true, regardless of moat_score."
    )

    resp = requests.post(
        "https://api.anthropic.com/v1/messages",
        headers={
            "x-api-key": ANTHROPIC_KEY,
            "anthropic-version": "2023-06-01",
            "content-type": "application/json",
        },
        json={
            "model": CLAUDE_MODEL,
            "max_tokens": 4000,
            "messages": [{"role": "user", "content": prompt}],
        },
        timeout=60,
    )
    resp.raise_for_status()
    text = resp.json()["content"][0]["text"].strip()
    text = re.sub(r"^```(json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        print("    ⚠ Could not parse Claude's response as JSON:")
        print(text[:500])
        return []


MAX_FUNDING_USD = 1_500_000  # hard cap when we actually have a real reported amount (tech.eu / Wamda)


def main():
    cache = load_cache()

    # Hard-reject over-cap funding before spending a Claude call on it — this is
    # real reported data (tech.eu / Wamda headlines), not a guess, so no need for
    # an LLM's judgment here.
    for c in cache:
        if not c.get("scored") and c.get("funding_usd") and c["funding_usd"] > MAX_FUNDING_USD:
            c["scored"] = True
            c["pass"] = False
            c["moat_score"] = 0
            c["reasoning"] = f"Funding ${c['funding_usd']:,.0f} exceeds the ${MAX_FUNDING_USD:,.0f} cap."

    unscored = [c for c in cache if not c.get("scored")]

    print(f"Unscored candidates in cache: {len(unscored)}")

    if unscored:
        print(f"Scoring {len(unscored)} candidates with Claude…")
        results = score_candidates(unscored)
        results_by_name = {str(r.get("name", "")).strip().lower(): r for r in results}

        for c in unscored:
            r = results_by_name.get(c["name"].strip().lower())
            if r:
                c["pass"] = r.get("pass", False)
                c["moat_score"] = r.get("moat_score", 0)
                c["geography_guess"] = r.get("geography_guess", "")
                c["reasoning"] = r.get("reasoning", "")
                if r.get("company_name"):
                    c["company_name"] = r["company_name"]
            c["scored"] = True

        save_cache(cache)

    candidates_for_sheet = [
        c for c in cache
        if c.get("scored") and not c.get("added_to_sheet")
        and c.get("pass") and c.get("moat_score", 0) >= MIN_MOAT_SCORE
    ]
    candidates_for_sheet.sort(key=lambda x: x.get("moat_score", 0), reverse=True)
    top = candidates_for_sheet[:MAX_NEW_LEADS]

    print(f"Qualifying (score >= {MIN_MOAT_SCORE}, not yet added): {len(candidates_for_sheet)}  |  Adding: {len(top)}")

    if not top:
        print("No qualifying leads this week — nothing added (quality bar not met).")
        return

    wb = openpyxl.load_workbook(FILE_PATH)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers) if h}

    for label in ("Source", "AI Moat Score", "AI Notes"):
        if label not in col:
            col[label] = ws.max_column + 1
            ws.cell(1, col[label]).value = label

    row_idx = ws.max_row + 1
    for item in top:
        ws.cell(row_idx, col["Organization Name"]).value = item.get("company_name") or item["name"]
        ws.cell(row_idx, col["Description"]).value = item.get("tagline", "") or item.get("description", "")[:200]
        ws.cell(row_idx, col["Website"]).value = item.get("website", "").split("?")[0].rstrip("/")
        if "Move to Pipeline?" in col:
            ws.cell(row_idx, col["Move to Pipeline?"]).value = False
        ws.cell(row_idx, col["Headquarters Location"]).value = item.get("geography_guess", "")
        ws.cell(row_idx, col["Source"]).value = item.get("source", "BetaList")
        ws.cell(row_idx, col["AI Moat Score"]).value = item.get("moat_score", "")
        notes = item.get("reasoning", "")
        if item.get("funding_usd"):
            notes += f" | Reported funding: ${item['funding_usd']:,.0f}"
        ws.cell(row_idx, col["AI Notes"]).value = notes
        print(f"  + {item['name']} — moat {item.get('moat_score')} — {item.get('reasoning','')}")
        row_idx += 1
        item["added_to_sheet"] = True

    wb.save(FILE_PATH)
    save_cache(cache)
    print(f"\nSaved {len(top)} new leads to: {FILE_PATH}")


if __name__ == "__main__":
    main()
