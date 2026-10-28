"""
Apollo enrichment for corporate lead lists (e.g. 'Madrid in Game - Lead corporate gaming'): one decision maker
per company, in the country of the local office, never a random employee of a multinational.

  python3 enrich_corporate_contacts.py <job_dir> [--max-credits N] [--dry]

Reads <job_dir>/original.xlsx (never modified), writes <job_dir>/enriched.xlsx. Expected columns (header row
found by 'Azienda'): Azienda, Gruppo / proprietà, Sede di riferimento, Contatto noto (pubblico),
Ruolo da invitare, Email (da PitchBook), Sito web. Only rows with an empty email are touched.

Per row:
  1. Domain from 'Sito web'; none -> Apollo company search by name (1 credit), exact name match only.
  2. Named contact ('Name, role...') -> people/match on name + domain (1 credit if an email comes back).
  3. Otherwise: free people search on the domain in the office's country, keep only decision-maker titles
     (country/general manager, MD, studio head/director, site lead, C-level, founder, head of / director /
     VP of partnerships, BD, marketing, esports, communication), titles echoing 'Ruolo da invitare' first;
     reveal the best (max 2 tries, 1 credit each when an email comes back).
  The email must be on the website domain or on the person's own company domain in Apollo.
Writes the email in 'Email (da PitchBook)' and adds 'Contatto trovato', 'Ruolo contatto', 'Fonte email'.
Answers cached in <job_dir>/apollo_cache.json (reruns cost nothing); stops at --max-credits.
"""
import json, os, re, sys, time
import openpyxl
import requests

JOB = sys.argv[1]
MAX_CREDITS = int(sys.argv[sys.argv.index("--max-credits") + 1]) if "--max-credits" in sys.argv else 80
DRY = "--dry" in sys.argv
ROOT = os.path.dirname(os.path.abspath(__file__))
KEY = next(l.split("=", 1)[1].strip().strip('"\'') for l in open(os.path.join(ROOT, ".env"))
           if l.startswith("APOLLO_KEY="))
HEADERS = {"x-api-key": KEY, "Content-Type": "application/json", "Accept": "application/json"}
CACHE_PATH = os.path.join(JOB, "apollo_cache.json")
cache = json.load(open(CACHE_PATH)) if os.path.exists(CACHE_PATH) else {}
credits = sum(1 for v in cache.values() if v.get("_credit"))

COUNTRY = [  # office city (as written in the sheet, Italian) -> Apollo location
    (r"parigi|boulogne|lione|lyon|montpellier|bordeaux|francia", "France"),
    (r"milano|roma|italia", "Italy"),
    (r"losanna|ginevra|zurigo|svizzera", "Switzerland"),
    (r"lisbona|oeiras|porto|portogallo", "Portugal"),
]
TOP = r"country manager|general manager|managing director|studio (head|director|manager|lead)|head of studio|" \
      r"site lead|site director|\bceo\b|chief executive|founder|president|presidente|director general|\bcountry\b"
MID = r"head of|director|\bvp\b|vice president|partnership|business development|marketing|esports|e-sports|" \
      r"communication|comunicaci|\bcmo\b|\bcoo\b|\bcto\b|\bcfo\b|chief"


def country_of(sede):
    # The first city listed is the office we invite ("Barcellona; HQ Parigi" -> Spain).
    s = re.split(r"[;+]| e | and ", (sede or "").lower())[0]
    return next((c for pat, c in COUNTRY if re.search(pat, s)), "Spain")


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def domain_of(v):
    s = re.sub(r"^https?://", "", str(v or "").strip(), flags=re.I)
    s = re.sub(r"^www\.", "", s, flags=re.I)
    return s.split("/")[0].split("?")[0].strip().lower()


def clean_email(e):
    e = (e or "").strip()
    return "" if not e or re.search(r"not_unlocked|@domain\.com$", e, re.I) else e


def email_ok(email, site_domain, person, names):
    # The person's own company domain counts only if that company is this one (or its group):
    # Apollo can return someone who has since moved to another company.
    host = email.split("@")[-1].lower()
    org = person.get("organization") or {}
    org_domain = (org.get("primary_domain") or "").lower()
    org_name = norm(org.get("name"))
    same_company = bool(org_name) and any(n and (n in org_name or org_name in n) for n in map(norm, names))
    return any(d and (host == d or host.endswith("." + d))
               for d in (site_domain, org_domain if same_company else ""))


def score(title, role_words):
    t = (title or "").lower()
    base = 0 if re.search(TOP, t) else 1 if re.search(MID, t) else None
    if base is None:
        return None  # not a decision maker: never revealed
    echo = sum(1 for w in role_words if w in t)
    return base * 10 - echo


def apollo(url, payload, costs_credit):
    global credits
    key = url + json.dumps(payload, sort_keys=True)
    if key in cache:
        return cache[key]
    if costs_credit and credits >= MAX_CREDITS:
        raise SystemExit(f"Credit cap {MAX_CREDITS} reached - stopping (progress cached).")
    if DRY and costs_credit:
        return {}
    for attempt in range(3):
        r = requests.post(url, json=payload, headers=HEADERS, timeout=30)
        if r.status_code != 429:
            break
        time.sleep(60 * (attempt + 1))
    if r.status_code != 200:
        raise SystemExit(f"Apollo {r.status_code} on {url}: {r.text[:200]} (progress cached)")
    data = r.json()
    data["_credit"] = costs_credit and bool(
        (data.get("person") or {}).get("email") or data.get("organizations") or data.get("accounts"))
    credits += bool(data["_credit"])
    cache[key] = data
    json.dump(cache, open(CACHE_PATH, "w"))
    time.sleep(1)
    return data


def main():
    wb = openpyxl.load_workbook(os.path.join(JOB, "original.xlsx"))
    ws = wb.worksheets[0]
    hr = next(r for r in range(1, 20) if str(ws.cell(r, 2).value or "").strip() == "Azienda")
    col = {str(ws.cell(hr, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
    C_CO, C_GRP, C_SEDE, C_KNOWN, C_ROLE, C_EMAIL, C_SITE = (col[k] for k in (
        "Azienda", "Gruppo / proprietà", "Sede di riferimento", "Contatto noto (pubblico)",
        "Ruolo da invitare", "Email (da PitchBook)", "Sito web"))
    C_NAME, C_TITLE, C_SRC = ws.max_column + 1, ws.max_column + 2, ws.max_column + 3
    for c, h in ((C_NAME, "Contatto trovato"), (C_TITLE, "Ruolo contatto"), (C_SRC, "Fonte email")):
        ws.cell(hr, c).value = h
        ws.cell(hr, c)._style = ws.cell(hr, C_EMAIL)._style
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 30
    stats = {}

    for r in range(hr + 1, ws.max_row + 1):
        company = str(ws.cell(r, C_CO).value or "").strip()
        if not company or not str(ws.cell(r, 1).value or "").strip():
            continue  # blank rows and the totals block
        if str(ws.cell(r, C_EMAIL).value or "").strip():
            continue

        def done(src, person=None):
            ws.cell(r, C_SRC).value = src
            if person:
                ws.cell(r, C_EMAIL).value = person["email"]
                ws.cell(r, C_NAME).value = person.get("name") or \
                    f"{person.get('first_name', '')} {person.get('last_name', '')}".strip()
                ws.cell(r, C_TITLE).value = person.get("title") or ""
            k = "found" if person else "not found"
            stats[k] = stats.get(k, 0) + 1
            print(f"row {r}: {company} -> {src}")

        domain, note = domain_of(ws.cell(r, C_SITE).value), ""
        if not domain:
            org = apollo("https://api.apollo.io/api/v1/mixed_companies/search",
                         {"q_organization_name": company, "per_page": 5}, True)
            hit = next((o for o in (org.get("organizations") or org.get("accounts") or [])
                        if norm(o.get("name")) == norm(company) and o.get("primary_domain")), None)
            if not hit:
                done("Not found: no website and company not found in Apollo")
                continue
            domain, note = hit["primary_domain"].lower(), f" (website {hit['primary_domain']} from Apollo)"
        country = country_of(ws.cell(r, C_SEDE).value)
        role_text = str(ws.cell(r, C_ROLE).value or "").lower()
        role_words = [w for w in re.findall(r"[a-z]{4,}", role_text)
                      if w not in {"head", "lead", "with", "della", "delle", "spain", "españa", "iberia"}]

        known = str(ws.cell(r, C_KNOWN).value or "").split(",")[0].strip()
        known = re.sub(r"\(.*?\)", "", known).strip()
        if known and len(known.split()) >= 2 and len(known) < 40:
            first, last = known.split(" ", 1)
            p = apollo("https://api.apollo.io/api/v1/people/match",
                       {"first_name": first, "last_name": last, "domain": domain}, True).get("person") or {}
            if clean_email(p.get("email")) and email_ok(p["email"], domain, p, (company, str(ws.cell(r, C_GRP).value or ""))):
                done("Apollo: named contact in the sheet" + note, p)
                continue

        people = apollo("https://api.apollo.io/api/v1/mixed_people/api_search",
                        {"q_organization_domains_list": [domain], "person_locations": [country],
                         "per_page": 100}, False).get("people") or []
        ranked = sorted(((score(p.get("title"), role_words), p) for p in people if p.get("id")),
                        key=lambda x: (x[0] is None, x[0] if x[0] is not None else 0,
                                       not x[1].get("has_email")))
        ranked = [p for s, p in ranked if s is not None]
        person = None
        for cand in ranked[:2]:
            p = apollo("https://api.apollo.io/api/v1/people/match", {"id": cand["id"]}, True).get("person") or {}
            if clean_email(p.get("email")) and email_ok(p["email"], domain, p, (company, str(ws.cell(r, C_GRP).value or ""))):
                person = p
                break
        if person:
            why = f"Apollo: decision maker in {country}" + (f" (named contact {known} not found)" if known else "")
            done(why + note, person)
        else:
            done(f"Not found: no decision maker with email on {domain} in {country}" + note)

    if not DRY:
        wb.save(os.path.join(JOB, "enriched.xlsx"))
    print(json.dumps(stats), "| credits used (incl. earlier runs):", credits, "/", MAX_CREDITS)


if __name__ == "__main__":
    main()
