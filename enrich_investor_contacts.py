"""
Apollo enrichment for investor lists (PitchBook-style exports): fills the missing contact email per firm.

  python3 enrich_investor_contacts.py <job_dir> [--max-credits N] [--dry]

<job_dir>/original.xlsx is read (never modified); the result goes to <job_dir>/enriched.xlsx with every
original sheet, formula and format kept. Only rows whose 'Primary Contact Email' is empty are touched.

Per row:
  1. Domain from 'Website' (HYPERLINK formulas included). No website -> Apollo company search by name
     (1 credit), domain used only if the company name matches exactly.
  2. A named 'Primary Contact' -> people/match on name + domain (1 credit when an email comes back).
  3. Otherwise, or if that person has no email: free people search on the domain, ranked
     Managing/General/Founding Partner > Partner > Founder/CEO > Principal/Director/Investment roles,
     then reveal the best one (1 credit).
  Emails must be on the firm's domain; placeholders are discarded.
Filled: 'Primary Contact Email' (+ name/title when the person differs or was missing) and a new
'Enrich source' column saying where each value came from. Every Apollo answer is cached in
<job_dir>/apollo_cache.json, so a rerun never spends credits twice. Stops at --max-credits.
The job dir holds personal data: keep it in one-off-sends/ (gitignored).
"""
import json, os, re, sys, time
import openpyxl
import requests

JOB = sys.argv[1]
MAX_CREDITS = int(sys.argv[sys.argv.index("--max-credits") + 1]) if "--max-credits" in sys.argv else 90
DRY = "--dry" in sys.argv
ROOT = os.path.dirname(os.path.abspath(__file__))
KEY = next(l.split("=", 1)[1].strip().strip('"\'') for l in open(os.path.join(ROOT, ".env"))
           if l.startswith("APOLLO_KEY="))
HEADERS = {"x-api-key": KEY, "Content-Type": "application/json", "Accept": "application/json"}
CACHE_PATH = os.path.join(JOB, "apollo_cache.json")
cache = json.load(open(CACHE_PATH)) if os.path.exists(CACHE_PATH) else {}
credits = sum(1 for k, v in cache.items() if v.get("_credit"))

RANKS = [
    ["managing partner", "general partner", "founding partner", "senior partner", "managing director"],
    ["partner"],
    ["founder", "ceo", "chief executive", "president"],
    ["principal", "director", "investment manager", "head of", "investor", "vp", "vice president"],
]


def rank(title):
    t = (title or "").lower()
    for i, keys in enumerate(RANKS):
        if any(k in t for k in keys):
            return i
    return 99


def norm(s):
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def domain_of(v):
    s = str(v or "").strip()
    m = re.search(r'HYPERLINK\("([^"]+)"', s, re.I)
    if m:
        s = m.group(1)
    s = re.sub(r"^https?://", "", s, flags=re.I)
    s = re.sub(r"^www\.", "", s, flags=re.I)
    return s.split("/")[0].split("?")[0].strip().lower()


def on_domain(email, domain):
    host = email.split("@")[-1].lower()
    return bool(domain) and (host == domain or host.endswith("." + domain))


def clean_email(e):
    e = (e or "").strip()
    return "" if not e or re.search(r"not_unlocked|@domain\.com$", e, re.I) else e


def apollo(url, payload, costs_credit):
    global credits
    key = url + json.dumps(payload, sort_keys=True)
    if key in cache:
        return cache[key]
    if costs_credit and credits >= MAX_CREDITS:
        raise SystemExit(f"Credit cap {MAX_CREDITS} reached - stopping (progress saved).")
    if DRY and costs_credit:
        return {"_dry": True}
    for attempt in range(3):
        r = requests.post(url, json=payload, headers=HEADERS, timeout=30)
        if r.status_code == 429:
            time.sleep(60 * (attempt + 1))
            continue
        break
    if r.status_code != 200:
        raise SystemExit(f"Apollo {r.status_code} on {url}: {r.text[:200]} (progress saved)")
    data = r.json()
    data["_credit"] = costs_credit and bool(
        (data.get("person") or {}).get("email") or data.get("organizations") or data.get("accounts"))
    credits += bool(data["_credit"])
    cache[key] = data
    json.dump(cache, open(CACHE_PATH, "w"))
    time.sleep(1)
    return data


def main():
    wb = openpyxl.load_workbook(os.path.join(JOB, "original.xlsx"))
    ws = wb["Data"]
    hr = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Investor ID")
    col = {str(ws.cell(hr, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value}
    C_NAME, C_EMAIL, C_TITLE, C_SITE, C_FIRM = (col[k] for k in (
        "Primary Contact", "Primary Contact Email", "Primary Contact Title", "Website", "Investors"))
    C_SRC = ws.max_column + 1
    ws.cell(hr, C_SRC).value = "Enrich source"
    ws.cell(hr, C_SRC)._style = ws.cell(hr, C_FIRM)._style
    stats = {}

    def note(r, src):
        ws.cell(r, C_SRC).value = src
        k = src.split(" (")[0].split(":")[0]
        stats[k] = stats.get(k, 0) + 1

    for r in range(hr + 1, ws.max_row + 1):
        firm = str(ws.cell(r, C_FIRM).value or "").strip()
        if not firm:
            continue
        if str(ws.cell(r, C_EMAIL).value or "").strip():
            continue
        name = str(ws.cell(r, C_NAME).value or "").strip()
        domain, src_domain = domain_of(ws.cell(r, C_SITE).value), ""
        if not domain:
            org = apollo("https://api.apollo.io/api/v1/mixed_companies/search",
                         {"q_organization_name": firm, "per_page": 5}, True)
            hit = next((o for o in (org.get("organizations") or org.get("accounts") or [])
                        if norm(o.get("name")) == norm(firm) and o.get("primary_domain")), None)
            if not hit:
                note(r, "Not found: no website, firm not found in Apollo")
                continue
            domain, src_domain = hit["primary_domain"].lower(), f", website {hit['primary_domain']} from Apollo"

        person = {}
        if name and " " in name:
            first, last = name.split(" ", 1)
            p = (apollo("https://api.apollo.io/api/v1/people/match",
                        {"first_name": first, "last_name": last, "domain": domain}, True).get("person") or {})
            if clean_email(p.get("email")) and on_domain(p["email"], domain):
                person = p
        how = "PitchBook contact" if person else ""
        if not person:
            people = apollo("https://api.apollo.io/api/v1/mixed_people/api_search",
                            {"q_organization_domains_list": [domain], "per_page": 25}, False).get("people") or []
            people = [p for p in people if p.get("id")]
            ranked = sorted([p for p in people if p.get("has_email")] or people, key=lambda p: rank(p.get("title")))
            for cand in ranked[:2]:  # at most 2 reveals per firm
                p = apollo("https://api.apollo.io/api/v1/people/match", {"id": cand["id"]}, True).get("person") or {}
                if clean_email(p.get("email")) and on_domain(p["email"], domain):
                    person = p
                    break
            how = "Senior person at the firm" + (f" (PitchBook contact {name} not found)" if name else "")
        if not person:
            note(r, f"Not found: no email on {domain} in Apollo{src_domain}")
            continue
        full = person.get("name") or f"{person.get('first_name', '')} {person.get('last_name', '')}".strip()
        ws.cell(r, C_EMAIL).value = person["email"]
        if not name or norm(full) != norm(name):
            ws.cell(r, C_NAME).value = full
            ws.cell(r, C_TITLE).value = person.get("title") or ""
        elif not str(ws.cell(r, C_TITLE).value or "").strip():
            ws.cell(r, C_TITLE).value = person.get("title") or ""
        note(r, "Apollo: " + how + src_domain)
        print(f"row {r}: {firm} -> {how}")

    if not DRY:
        wb.save(os.path.join(JOB, "enriched.xlsx"))
    print(json.dumps(stats, indent=1), "\ncredits used (incl. earlier runs):", credits, "/", MAX_CREDITS)


if __name__ == "__main__":
    main()
